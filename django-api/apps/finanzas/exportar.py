"""
El Excel del mes para el contador: una hoja por cosa, con lo que necesita
para las declaraciones (ventas, compras con RUC, retenciones, nómina).

Los importes van como número, no como texto: el contador los suma.
"""

import io
from decimal import Decimal

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill

from apps.common.document_style import get_document_style
from apps.finanzas import calculos


def _hoja(wb, titulo, encabezado, filas, tema, primera=False, totales=None):
    ws = wb.active if primera else wb.create_sheet()
    ws.title = titulo[:31]
    negrita = Font(name=tema["font"], size=tema["size"], bold=True, color=tema["header_text"])
    relleno = PatternFill("solid", start_color=tema["header_bg"])
    cuerpo = Font(name=tema["font"], size=tema["size"], color=tema["body_text"])
    ws.append(encabezado)
    for celda in ws[1]:
        celda.font, celda.fill = negrita, relleno
        celda.alignment = Alignment(vertical="center")
    for fila in filas:
        ws.append([float(v) if isinstance(v, Decimal) else v for v in fila])
    if totales:
        ws.append([float(v) if isinstance(v, Decimal) else v for v in totales])
        for celda in ws[ws.max_row]:
            celda.font = Font(name=tema["font"], size=tema["size"], bold=True)
    for fila in ws.iter_rows(min_row=2, max_row=ws.max_row - (1 if totales else 0)):
        for celda in fila:
            celda.font = cuerpo
    for col in ws.columns:
        largo = max(len(str(c.value or "")) for c in col)
        ws.column_dimensions[col[0].column_letter].width = min(max(largo + 3, 10), 50)
        for c in col[1:]:
            if isinstance(c.value, float):
                c.number_format = '#,##0.00'
    ws.freeze_panes = "A2"
    return ws


def libro_del_mes(tenant, desde, hasta):
    from apps.finanzas.models import Egreso, Liquidacion, RolPagos

    tema = get_document_style(tenant).xlsx_theme()
    wb = Workbook()

    pagos = calculos.cobros(tenant, desde, hasta).select_related("patient", "profesional").order_by("date")
    filas = [[p.date.isoformat(), p.numero_comprobante or "", p.patient.full_name, p.patient.national_id,
              p.get_method_display(), p.profesional.full_name if p.profesional_id else "", p.amount]
             for p in pagos]
    _hoja(wb, "Cobros", ["Fecha", "Comprobante", "Paciente", "Cédula", "Forma de pago", "Profesional",
                         "Monto"], filas, tema, primera=True,
          totales=["Total", "", "", "", "", "", sum((p.amount for p in pagos), Decimal("0"))])

    egresos = Egreso.objects.filter(tenant=tenant, fecha__gte=desde, fecha__lte=hasta) \
        .select_related("categoria").order_by("fecha")
    filas = [[e.fecha.isoformat(), e.categoria.nombre, e.concepto, e.proveedor, e.proveedor_ruc,
              e.numero_factura, e.get_forma_pago_display(), e.get_estado_display(),
              "ANULADO: " + e.motivo_anulacion if e.anulado else "", e.monto]
             for e in egresos]
    validos = [e for e in egresos if not e.anulado]
    _hoja(wb, "Egresos", ["Fecha", "Categoría", "Concepto", "Proveedor", "RUC / Cédula", "N.º factura",
                          "Forma de pago", "Estado", "Observación", "Monto"], filas, tema,
          totales=["Total (sin anulados)", "", "", "", "", "", "", "", "",
                   sum((e.monto for e in validos), Decimal("0"))])

    liqs = Liquidacion.objects.filter(tenant=tenant, hasta__gte=desde, hasta__lte=hasta) \
        .exclude(estado=Liquidacion.Estado.ANULADA).select_related("profesional").order_by("numero")
    filas = [[f"{q.numero:05d}", q.profesional.full_name,
              getattr(getattr(q.profesional, "perfil_laboral", None), "cedula", ""),
              f"{q.desde.isoformat()} a {q.hasta.isoformat()}", q.base, q.honorario, q.descuentos,
              q.ajustes, q.retencion_porcentaje, q.retencion, q.a_pagar, q.get_estado_display()]
             for q in liqs]
    _hoja(wb, "Honorarios y retenciones", ["N.º", "Profesional", "Cédula / RUC", "Período",
                                           "Cobrado", "Honorario", "Descuentos", "Ajustes",
                                           "% Retención IR", "Retención IR", "A pagar", "Estado"],
          filas, tema)

    roles = RolPagos.objects.filter(tenant=tenant, mes=desde).exclude(
        estado=RolPagos.Estado.ANULADO).select_related("empleado")
    filas = []
    for r in roles:
        c = r.calculo or {}
        aporte = next((d["valor"] for d in c.get("descuentos", []) if d["concepto"].startswith("Aporte")), "0")
        filas.append([r.empleado.full_name, getattr(getattr(r.empleado, "perfil_laboral", None), "cedula", ""),
                      r.dias_trabajados, Decimal(c.get("total_ingresos", "0")), Decimal(aporte),
                      Decimal(c.get("total_descuentos", "0")), r.neto, Decimal(c.get("aporte_patronal", "0")),
                      Decimal(c.get("total_provisiones", "0")), r.costo_empleador, r.get_estado_display()])
    _hoja(wb, "Rol de pagos", ["Trabajador", "Cédula", "Días", "Total ingresos", "Aporte personal IESS",
                               "Total descuentos", "Neto a recibir", "Aporte patronal", "Provisiones",
                               "Costo para la clínica", "Estado"], filas, tema)

    r = calculos.resumen(tenant, desde, hasta)
    resumen = wb.create_sheet("Resumen", 0)
    for fila in (["Mes", desde.strftime("%Y-%m")], ["Ingresos cobrados", float(r["ingresos"])],
                 ["Egresos pagados", float(r["egresos"])], ["Utilidad", float(r["utilidad"])],
                 ["Por cobrar a pacientes", float(r["por_cobrar"])], ["Por pagar", float(r["por_pagar"])]):
        resumen.append(fila)
    # Lo que se debe al IESS y al SRI por este mes, y si ya se pagó.
    for o in calculos.obligaciones(tenant, desde):
        resumen.append([f"{o['nombre']} ({'pagada' if o['pagado'] else 'vence ' + o['vence']})",
                        float(o["monto"])])
    resumen.column_dimensions["A"].width = 28
    resumen.column_dimensions["B"].width = 16
    for c in resumen["B"][1:]:
        c.number_format = '#,##0.00'
    for c in resumen["A"]:
        c.font = Font(bold=True)

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
