"""Finanzas de la clínica: cuentas, permisos y papeles."""

import io
from datetime import date, timedelta
from decimal import Decimal

from django.core.cache import cache
from django.test import TestCase
from django.utils import timezone
from openpyxl import load_workbook
from rest_framework.test import APITestCase

from apps.accounts.models import AuditLog, User
from apps.agenda.models import Appointment, Doctor
from apps.billing.models import Payment
from apps.common.models import Tenant
from apps.finanzas import calculos
from apps.finanzas.documentos import en_letras
from apps.finanzas.models import (
    CierreCaja,
    ConfiguracionFinanzas,
    Egreso,
    PerfilLaboral,
    categoria,
)
from apps.patients.models import Patient

D = Decimal
CLAVE = "clave-larga-8gatos"


class Base(APITestCase):
    def setUp(self):
        cache.clear()
        self.addCleanup(cache.clear)
        # Finanzas nace «en prueba»: solo la ven las clínicas de prueba.
        self.tenant = Tenant.objects.create(name="Clínica Finanzas", es_prueba=True)
        self.admin = User.objects.create_user(email="admin@fin.ec", password=CLAVE, role="admin",
                                              tenant=self.tenant, full_name="Ana Admin")
        self.doc_user = User.objects.create_user(email="doc@fin.ec", password=CLAVE, role="doctor",
                                                 tenant=self.tenant, full_name="Diego Doctor")
        self.doctor = Doctor.objects.create(tenant=self.tenant, user=self.doc_user)
        self.recepcion = User.objects.create_user(email="rec@fin.ec", password=CLAVE, role="reception",
                                                  tenant=self.tenant, full_name="Rosa Recepción")
        self.paciente = Patient.objects.create(tenant=self.tenant, first_name="Pedro", last_name="Paz",
                                               national_id="1710034065", phone="0991111111")
        self.mes = date(2026, 9, 1)
        self.client.force_authenticate(self.admin)

    def cobro(self, monto, dia=None, metodo="cash", profesional="doc"):
        return Payment.objects.create(
            tenant=self.tenant, patient=self.paciente, amount=D(monto), method=metodo,
            date=dia or self.mes.replace(day=10),
            profesional=self.doc_user if profesional == "doc" else profesional)

    def egreso(self, monto, clave="otros", dia=None, **extra):
        return Egreso.objects.create(tenant=self.tenant, fecha=dia or self.mes.replace(day=12),
                                     categoria=categoria(self.tenant, clave), concepto="Gasto",
                                     monto=D(monto), **extra)


class LiquidacionCalculoTests(Base):
    def test_porcentaje_de_lo_cobrado_menos_laboratorio_y_retencion(self):
        PerfilLaboral.objects.create(tenant=self.tenant, usuario=self.doc_user, modalidad="honorarios",
                                     porcentaje=D("40"), emite_factura=True)
        self.cobro("300")
        self.cobro("200", metodo="transfer")
        self.cobro("999", profesional=None)                     # no es suyo
        self.cobro("50", dia=date(2026, 10, 2))                 # otro mes
        self.egreso("60", "laboratorio", profesional=self.doc_user, descontar_al_profesional=True)
        self.egreso("70", "laboratorio", profesional=self.doc_user)   # no se le descuenta
        r = calculos.liquidacion(self.tenant, self.doc_user, date(2026, 9, 1), date(2026, 9, 30),
                                 ajustes=D("-20"))
        self.assertEqual(r["base"], D("500.00"))
        self.assertEqual(r["honorario"], D("200.00"))
        self.assertEqual(r["descuentos"], D("60.00"))
        self.assertEqual(r["retencion"], D("20.00"))            # 10 % de 200: el anticipo no cambia la factura
        self.assertEqual(r["a_pagar"], D("100.00"))             # 200 − 60 − 20 − 20
        self.assertEqual(len(r["detalle"]["cobros"]), 2)

    def test_sin_factura_no_hay_retencion_y_el_bono_suma(self):
        PerfilLaboral.objects.create(tenant=self.tenant, usuario=self.doc_user, modalidad="honorarios",
                                     porcentaje=D("50"), emite_factura=False)
        self.cobro("100")
        r = calculos.liquidacion(self.tenant, self.doc_user, date(2026, 9, 1), date(2026, 9, 30), D("15"))
        self.assertEqual((r["retencion"], r["a_pagar"]), (D("0.00"), D("65.00")))

    def test_monto_fijo(self):
        PerfilLaboral.objects.create(tenant=self.tenant, usuario=self.doc_user, modalidad="honorarios",
                                     base_honorarios="fijo", monto_fijo=D("800"), emite_factura=True)
        self.cobro("5000")
        r = calculos.liquidacion(self.tenant, self.doc_user, date(2026, 9, 1), date(2026, 9, 30))
        self.assertEqual((r["honorario"], r["retencion"], r["a_pagar"]), (D("800.00"), D("80.00"), D("720.00")))

    def test_la_retencion_sale_de_la_configuracion(self):
        config = ConfiguracionFinanzas.de(self.tenant)
        config.retencion_honorarios = D("8")
        config.save()
        PerfilLaboral.objects.create(tenant=self.tenant, usuario=self.doc_user, porcentaje=D("100"))
        self.cobro("100")
        r = calculos.liquidacion(self.tenant, self.doc_user, date(2026, 9, 1), date(2026, 9, 30))
        self.assertEqual(r["retencion"], D("8.00"))


class RolCalculoTests(Base):
    def perfil(self, **kw):
        datos = {"modalidad": "dependencia", "sueldo": D("600"), "fecha_ingreso": date(2026, 1, 15),
                 "fondos_reserva_mensual": True}
        datos.update(kw)
        return PerfilLaboral.objects.create(tenant=self.tenant, usuario=self.recepcion, **datos)

    def test_primer_anio_sin_fondos_y_decimos_provisionados(self):
        self.perfil()
        r = calculos.rol(self.tenant, self.recepcion, self.mes)
        ingresos = {f["concepto"]: D(f["valor"]) for f in r["ingresos"]}
        self.assertEqual(ingresos, {"Sueldo": D("600.00")})
        self.assertEqual(r["descuentos"][0]["valor"], "56.70")           # 9,45 %
        self.assertEqual(r["neto"], "543.30")
        self.assertEqual(r["aporte_patronal"], "72.90")                  # 12,15 %
        provisiones = {f["concepto"]: f["valor"] for f in r["provisiones"]}
        self.assertEqual(provisiones["Décimo tercer sueldo"], "50.00")
        self.assertEqual(provisiones["Décimo cuarto sueldo"], "40.17")   # 482 / 12
        self.assertNotIn("Fondos de reserva", provisiones)
        self.assertIsNotNone(r["fondos_reserva_desde"])

    def test_con_un_anio_y_todo_mensualizado(self):
        self.perfil(fecha_ingreso=date(2025, 3, 1), decimo_tercero_mensual=True, decimo_cuarto_mensual=True)
        r = calculos.rol(self.tenant, self.recepcion, self.mes, horas_extra=D("40"), anticipos=D("100"))
        ingresos = {f["concepto"]: f["valor"] for f in r["ingresos"]}
        self.assertEqual(ingresos["Otras horas extra"], "40.00")
        self.assertEqual(ingresos["Décimo tercer sueldo"], "53.33")      # (600 + 40) / 12
        self.assertEqual(ingresos["Fondos de reserva"], "53.31")          # 8,33 % de 640
        self.assertEqual(r["descuentos"][0]["valor"], "60.48")           # 9,45 % de 640
        total = D("640") + D("53.33") + D("40.17") + D("53.31")
        self.assertEqual(r["total_ingresos"], str(total))
        self.assertEqual(r["neto"], str(total - D("60.48") - D("100")))

    def test_dias_trabajados_proporcionales(self):
        self.perfil()
        r = calculos.rol(self.tenant, self.recepcion, self.mes, dias_trabajados=15)
        self.assertEqual(r["ingresos"][0], {"concepto": "Sueldo (15 días)", "valor": "300.00"})

    def test_el_salario_basico_es_ajustable(self):
        config = ConfiguracionFinanzas.de(self.tenant)
        config.salario_basico = D("500")
        config.save()
        self.perfil()
        provisiones = {f["concepto"]: f["valor"] for f in calculos.rol(self.tenant, self.recepcion, self.mes)["provisiones"]}
        self.assertEqual(provisiones["Décimo cuarto sueldo"], "41.67")


class ResumenYCajaTests(Base):
    def test_utilidad_del_mes_sin_anulados_ni_pendientes(self):
        self.cobro("500")
        self.cobro("100", metodo="card")
        self.egreso("200", "arriendo")
        self.egreso("50", anulado=True)
        self.egreso("80", estado=Egreso.Estado.POR_PAGAR, vence=self.mes.replace(day=20))
        r = calculos.resumen(self.tenant, date(2026, 9, 1), date(2026, 9, 30))
        self.assertEqual((r["ingresos"], r["egresos"], r["utilidad"]), ("600.00", "200.00", "400.00"))
        self.assertEqual(r["por_pagar"], "80.00")
        self.assertEqual(r["margen"], "66.67")

    def test_la_caja_arranca_con_lo_que_quedo_del_cierre_anterior(self):
        dia = self.mes.replace(day=10)
        CierreCaja.objects.create(tenant=self.tenant, fecha=dia - timedelta(days=1), contado=D("150"),
                                  retirado=D("100"))
        self.cobro("80", dia=dia)
        self.cobro("40", dia=dia, metodo="transfer")
        self.egreso("30", dia=dia)
        c = calculos.caja(self.tenant, dia)
        self.assertEqual((c["saldo_inicial"], c["cobros_efectivo"], c["egresos_efectivo"], c["esperado"]),
                         ("50.00", "80.00", "30.00", "100.00"))
        self.assertEqual(c["otros_cobros"], "40.00")


class ApiTests(Base):
    def test_sin_la_funcion_no_se_entra(self):
        self.client.force_authenticate(self.recepcion)
        self.assertEqual(self.client.get("/api/v1/finanzas/resumen/").status_code, 403)
        self.recepcion.funciones = {"finanzas": True}
        self.recepcion.save()
        self.assertEqual(self.client.get("/api/v1/finanzas/resumen/").status_code, 200)

    def test_en_una_clinica_que_no_es_de_prueba_aun_no_existe(self):
        self.tenant.es_prueba = False
        self.tenant.save()
        self.assertEqual(self.client.get("/api/v1/finanzas/resumen/").status_code, 403)

    def test_resumen_con_serie_y_mes_anterior(self):
        self.cobro("100")
        r = self.client.get("/api/v1/finanzas/resumen/?mes=2026-09")
        self.assertEqual(r.status_code, 200, r.data)
        self.assertEqual(r.data["ingresos"], "100.00")
        self.assertEqual(len(r.data["serie"]), 6)
        self.assertEqual(r.data["serie"][-1]["mes"], "2026-09")

    def test_egreso_con_ruc_mal_escrito_y_anulacion_con_motivo(self):
        cat = categoria(self.tenant, "arriendo")
        r = self.client.post("/api/v1/finanzas/egresos/", {
            "fecha": "2026-09-05", "categoria": str(cat.id), "concepto": "Arriendo septiembre",
            "monto": "450", "proveedor_ruc": "123"}, format="json")
        self.assertEqual(r.status_code, 400)
        r = self.client.post("/api/v1/finanzas/egresos/", {
            "fecha": "2026-09-05", "categoria": str(cat.id), "concepto": "Arriendo septiembre",
            "monto": "450", "proveedor_ruc": "1790012345001", "numero_factura": "001-001-000000123"},
            format="json")
        self.assertEqual(r.status_code, 201, r.data)
        url = f"/api/v1/finanzas/egresos/{r.data['id']}/anular/"
        self.assertEqual(self.client.post(url, {}, format="json").status_code, 400)
        r = self.client.post(url, {"motivo": "Duplicado"}, format="json")
        self.assertTrue(r.data["anulado"])
        self.assertTrue(AuditLog.objects.filter(action="void_expense").exists())

    def test_formulario_del_panel_con_factura_adjunta(self):
        """Como lo manda el panel: multipart, campos vacíos y la factura en PDF."""
        from django.core.files.uploadedfile import SimpleUploadedFile

        cat = categoria(self.tenant, "laboratorio")
        datos = {"fecha": "2026-09-08", "categoria": str(cat.id), "concepto": "Corona 36", "monto": "120",
                 "forma_pago": "efectivo", "estado": "por_pagar", "proveedor": "", "proveedor_ruc": "",
                 "numero_factura": "", "profesional": "", "descontar_al_profesional": "false",
                 "comprobante": SimpleUploadedFile("factura.pdf", b"%PDF-1.4 prueba", "application/pdf")}
        r = self.client.post("/api/v1/finanzas/egresos/", datos, format="multipart")
        self.assertEqual(r.status_code, 201, r.data)
        self.assertTrue(r.data["tiene_comprobante"])
        self.assertEqual(r.data["vence"], "2026-09-08")
        r = self.client.get(f"/api/v1/finanzas/egresos/{r.data['id']}/comprobante/")
        self.assertEqual(r.status_code, 200)
        self.assertEqual(b"".join(r.streaming_content), b"%PDF-1.4 prueba")
        malo = {**datos, "comprobante": SimpleUploadedFile("virus.exe", b"MZ", "application/octet-stream")}
        self.assertEqual(self.client.post("/api/v1/finanzas/egresos/", malo, format="multipart").status_code, 400)

    def test_liquidacion_completa_y_sin_solapes(self):
        PerfilLaboral.objects.create(tenant=self.tenant, usuario=self.doc_user, modalidad="honorarios",
                                     porcentaje=D("40"))
        self.cobro("500")
        datos = {"profesional": str(self.doc_user.id), "desde": "2026-09-01", "hasta": "2026-09-30"}
        previa = self.client.get("/api/v1/finanzas/liquidaciones/previa/", datos)
        self.assertEqual(previa.data["a_pagar"], "180.00")
        r = self.client.post("/api/v1/finanzas/liquidaciones/", datos, format="json")
        self.assertEqual((r.status_code, r.data["numero"], r.data["a_pagar"]), (201, 1, "180.00"))
        self.assertEqual(self.client.post("/api/v1/finanzas/liquidaciones/", {**datos, "desde": "2026-09-15"},
                                          format="json").status_code, 400)
        pdf = self.client.get(f"/api/v1/finanzas/liquidaciones/{r.data['id']}/pdf/")
        self.assertEqual(pdf["Content-Type"], "application/pdf")
        self.assertTrue(pdf.content.startswith(b"%PDF"))

        pagada = self.client.post(f"/api/v1/finanzas/liquidaciones/{r.data['id']}/pagar/", {}, format="json")
        self.assertEqual(pagada.data["estado"], "pagada")
        egreso = Egreso.objects.get(origen="liquidacion")
        self.assertEqual((egreso.monto, egreso.categoria.clave), (D("180.00"), "honorarios"))
        # Un egreso de una liquidación no se toca a mano.
        self.assertEqual(self.client.post(f"/api/v1/finanzas/egresos/{egreso.id}/anular/",
                                          {"motivo": "x"}, format="json").status_code, 400)
        self.client.post(f"/api/v1/finanzas/liquidaciones/{r.data['id']}/anular/", {}, format="json")
        egreso.refresh_from_db()
        self.assertTrue(egreso.anulado)

    def test_un_cobro_ya_liquidado_no_cambia_de_profesional(self):
        PerfilLaboral.objects.create(tenant=self.tenant, usuario=self.doc_user, porcentaje=D("40"))
        pago = self.cobro("100")
        self.client.post("/api/v1/finanzas/liquidaciones/", {
            "profesional": str(self.doc_user.id), "desde": "2026-09-01", "hasta": "2026-09-30"}, format="json")
        r = self.client.patch(f"/api/v1/finanzas/cobros/{pago.id}/", {"profesional": str(self.admin.id)},
                              format="json")
        self.assertEqual(r.status_code, 400)
        self.assertIn("liquidación", r.data["detail"])

    def test_roles_del_mes_pagar_y_pdf(self):
        PerfilLaboral.objects.create(tenant=self.tenant, usuario=self.recepcion, modalidad="dependencia",
                                     sueldo=D("482"))
        r = self.client.post("/api/v1/finanzas/roles/generar/", {"mes": "2026-09"}, format="json")
        self.assertEqual(r.data["creados"], 1)
        self.assertEqual(self.client.post("/api/v1/finanzas/roles/generar/", {"mes": "2026-09"},
                                          format="json").data["creados"], 0)
        rol = self.client.get("/api/v1/finanzas/roles/?mes=2026-09").data[0]
        self.assertEqual(rol["neto"], "436.45")                          # 482 − 9,45 %
        # Se rehace con horas extra mientras está pendiente.
        r = self.client.post("/api/v1/finanzas/roles/", {"empleado": str(self.recepcion.id), "mes": "2026-09",
                                                         "horas_extra": "20"}, format="json")
        self.assertEqual(r.status_code, 200, r.data)
        self.assertEqual(r.data["neto"], "454.56")
        self.client.post(f"/api/v1/finanzas/roles/{rol['id']}/pagar/", {}, format="json")
        self.assertTrue(Egreso.objects.filter(origen="rol", categoria__clave="sueldos",
                                              monto=D("454.56")).exists())
        pdf = self.client.get(f"/api/v1/finanzas/roles/{rol['id']}/pdf/")
        self.assertTrue(pdf.content.startswith(b"%PDF"))

    def test_cierre_de_caja(self):
        hoy = timezone.localdate()
        self.cobro("120", dia=hoy)
        r = self.client.post("/api/v1/finanzas/caja/", {"fecha": hoy.isoformat(), "contado": "115",
                                                        "retirado": "100"}, format="json")
        self.assertEqual(r.status_code, 201, r.data)
        self.assertEqual((r.data["esperado"], r.data["diferencia"]), ("120.00", "-5.00"))
        mañana = (hoy + timedelta(days=1)).isoformat()
        self.assertEqual(self.client.post("/api/v1/finanzas/caja/", {"fecha": mañana, "contado": "1"},
                                          format="json").status_code, 400)

    def test_comprobante_de_cobro_numerado_y_estable(self):
        pago = self.cobro("35.50")
        otro = self.cobro("10")
        r1 = self.client.get(f"/api/v1/finanzas/cobros/{pago.id}/comprobante/")
        self.assertTrue(r1.content.startswith(b"%PDF"))
        self.client.get(f"/api/v1/finanzas/cobros/{otro.id}/comprobante/")
        self.client.get(f"/api/v1/finanzas/cobros/{pago.id}/comprobante/")
        pago.refresh_from_db()
        otro.refresh_from_db()
        self.assertEqual((pago.numero_comprobante, otro.numero_comprobante), (1, 2))

    def test_excel_para_el_contador(self):
        self.cobro("100")
        self.egreso("40", "insumos", proveedor="Dental Sur", proveedor_ruc="1790012345001")
        r = self.client.get("/api/v1/finanzas/exportar/?mes=2026-09")
        libro = load_workbook(io.BytesIO(r.content))
        self.assertEqual(libro.sheetnames, ["Resumen", "Cobros", "Egresos", "Honorarios y retenciones",
                                            "Rol de pagos"])
        self.assertEqual(libro["Resumen"]["B4"].value, 60.0)

    def test_equipo_y_perfil_ajustable(self):
        r = self.client.patch(f"/api/v1/finanzas/equipo/{self.doc_user.id}/",
                              {"modalidad": "honorarios", "porcentaje": "45", "cedula": "1710034065"},
                              format="json")
        self.assertEqual(r.status_code, 200, r.data)
        equipo = {f["nombre"]: f for f in self.client.get("/api/v1/finanzas/equipo/").data}
        self.assertEqual(equipo["Diego Doctor"]["porcentaje"], "45.00")
        self.assertEqual(self.client.patch(f"/api/v1/finanzas/equipo/{self.doc_user.id}/",
                                           {"porcentaje": "140"}, format="json").status_code, 400)


class CobroConProfesionalTests(Base):
    def test_al_cobrar_se_propone_el_profesional_de_su_ultima_cita(self):
        ahora = timezone.now()
        Appointment.objects.create(tenant=self.tenant, patient=self.paciente, doctor=self.doctor,
                                   scheduled_start=ahora - timedelta(days=2),
                                   scheduled_end=ahora - timedelta(days=2, hours=-1), status="completed")
        r = self.client.post(f"/api/v1/patients/{self.paciente.id}/payments/",
                             {"amount": "40", "method": "cash", "date": timezone.localdate().isoformat()},
                             format="json")
        self.assertEqual(r.status_code, 201, r.data)
        self.assertEqual(r.data["profesional"], str(self.doc_user.id))
        r = self.client.post(f"/api/v1/patients/{self.paciente.id}/payments/",
                             {"amount": "40", "method": "cash", "date": timezone.localdate().isoformat(),
                              "profesional": str(self.admin.id)}, format="json")
        self.assertEqual(r.status_code, 201, r.data)
        self.assertEqual(r.data["profesional_nombre"], "Ana Admin")


class CompraDeInventarioTests(Base):
    def test_una_compra_con_costo_queda_como_egreso_de_insumos(self):
        from apps.inventory.models import Product

        producto = Product.objects.create(tenant=self.tenant, name="Resina A2", unit="jeringa")
        r = self.client.post(f"/api/v1/products/{producto.id}/batches/", {
            "product": str(producto.id), "batch_number": "L-77", "quantity": "5", "costo_total": "62.50", "proveedor": "Dental Sur"},
            format="json")
        self.assertEqual(r.status_code, 201, r.data)
        egreso = Egreso.objects.get(origen="inventario")
        self.assertEqual((egreso.monto, egreso.categoria.clave), (D("62.50"), "insumos"))
        self.client.post(f"/api/v1/products/{producto.id}/batches/", {"product": str(producto.id), "batch_number": "L-78", "quantity": "1"},
                         format="json")
        self.assertEqual(Egreso.objects.count(), 1)


class EnLetrasTests(TestCase):
    def test_como_en_los_recibos(self):
        self.assertEqual(en_letras("121.50"), "Ciento veintiún con 50/100 dólares")
        self.assertEqual(en_letras("1001"), "Mil un con 00/100 dólares")
        self.assertEqual(en_letras("31"), "Treinta y un con 00/100 dólares")
